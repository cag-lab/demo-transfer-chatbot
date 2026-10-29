"""Core retrieval-augmented generation (RAG) logic.

Everything runs on your machine except two calls to YOUR Azure OpenAI resource:
  1. embeddings  (turn text into vectors for search)
  2. chat        (write the answer from the retrieved passages)

The search index is a pair of local files in ./index (no database server needed).
"""
from __future__ import annotations

import json
import os
import re
from dataclasses import dataclass, asdict
from pathlib import Path

import numpy as np
from dotenv import load_dotenv

load_dotenv()

ROOT = Path(__file__).parent
DOCS_DIR = Path(os.getenv("DOCS_DIR", ROOT / "docs"))
INDEX_DIR = Path(os.getenv("INDEX_DIR", ROOT / "index"))
TOP_K = int(os.getenv("TOP_K", "5"))
MIN_SCORE = float(os.getenv("MIN_SCORE", "0.25"))  # below this, a passage is treated as irrelevant

SYSTEM_PROMPT = """You are an assistant that answers questions about the institution's ICU transfer policies.

Rules:
- Answer ONLY from the numbered policy excerpts provided in the user message. Do not use outside medical knowledge.
- Cite every factual statement with the excerpt number in square brackets, e.g. [2].
- Quote numeric thresholds and times exactly as written in the policy.
- If the excerpts do not contain the answer, say: "I couldn't find that in the policy documents." Then suggest who to ask if the excerpts name a role.
- If the question is about a specific patient's care, remind the user that policies do not replace clinical judgment.
- Be concise. Use short bullet points when listing criteria.
"""


# --------------------------------------------------------------------------- #
# Azure OpenAI client
# --------------------------------------------------------------------------- #
_client = None


def get_client():
    """Create the Azure OpenAI client from environment variables (.env)."""
    global _client
    if _client is None:
        from openai import AzureOpenAI

        missing = [v for v in ("AZURE_OPENAI_ENDPOINT", "AZURE_OPENAI_API_KEY",
                               "AZURE_OPENAI_CHAT_DEPLOYMENT", "AZURE_OPENAI_EMBEDDING_DEPLOYMENT")
                   if not os.getenv(v)]
        if missing:
            raise RuntimeError(
                "Missing settings in .env: " + ", ".join(missing)
                + ". Copy .env.example to .env and fill it in."
            )
        _client = AzureOpenAI(
            azure_endpoint=os.environ["AZURE_OPENAI_ENDPOINT"],
            api_key=os.environ["AZURE_OPENAI_API_KEY"],
            api_version=os.getenv("AZURE_OPENAI_API_VERSION", "2024-10-21"),
        )
    return _client


def set_client(client) -> None:
    """Swap in a different client (used by tests)."""
    global _client
    _client = client


# --------------------------------------------------------------------------- #
# Loading and chunking documents
# --------------------------------------------------------------------------- #
@dataclass
class Chunk:
    id: str
    source: str     # file name
    title: str      # document title (first H1)
    section: str    # nearest heading
    text: str


def read_document(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in (".md", ".txt"):
        return path.read_text(encoding="utf-8", errors="ignore")
    if suffix == ".pdf":
        from pypdf import PdfReader
        return "\n\n".join(page.extract_text() or "" for page in PdfReader(str(path)).pages)
    if suffix == ".docx":
        import docx
        d = docx.Document(str(path))
        lines = []
        for p in d.paragraphs:
            style = (p.style.name or "").lower() if p.style is not None else ""
            if style.startswith("heading") and p.text.strip():
                level = re.sub(r"\D", "", style) or "2"
                lines.append("#" * int(level) + " " + p.text.strip())
            else:
                lines.append(p.text)
        for table in d.tables:
            for row in table.rows:
                lines.append(" | ".join(c.text.strip() for c in row.cells))
        return "\n".join(lines)
    raise ValueError(f"Unsupported file type: {path.name}")


def chunk_document(text: str, source: str, max_words: int = 350, overlap_words: int = 50) -> list[Chunk]:
    """Split on headings first, then split long sections by word count with overlap."""
    title_match = re.search(r"^#\s+(.+)$", text, flags=re.M)
    title = title_match.group(1).strip() if title_match else Path(source).stem

    sections: list[tuple[str, str]] = []
    current_heading, buf = "Introduction", []
    for line in text.splitlines():
        m = re.match(r"^#{1,4}\s+(.+)$", line)
        if m:
            if "".join(buf).strip():
                sections.append((current_heading, "\n".join(buf).strip()))
            current_heading, buf = m.group(1).strip(), []
        else:
            buf.append(line)
    if "".join(buf).strip():
        sections.append((current_heading, "\n".join(buf).strip()))

    chunks: list[Chunk] = []
    for heading, body in sections:
        words = body.split()
        start = 0
        while start < len(words):
            piece = " ".join(words[start:start + max_words])
            chunks.append(Chunk(
                id=f"{source}#{len(chunks)}", source=source, title=title,
                section=heading, text=piece,
            ))
            if start + max_words >= len(words):
                break
            start += max_words - overlap_words
    return chunks


def load_chunks(docs_dir: Path = DOCS_DIR) -> list[Chunk]:
    chunks: list[Chunk] = []
    for path in sorted(docs_dir.rglob("*")):
        if path.is_file() and path.suffix.lower() in (".md", ".txt", ".pdf", ".docx"):
            text = read_document(path)
            if text.strip():
                chunks.extend(chunk_document(text, path.name))
    return chunks


# --------------------------------------------------------------------------- #
# Embeddings and the local index
# --------------------------------------------------------------------------- #
def embed(texts: list[str], batch_size: int = 16) -> np.ndarray:
    client = get_client()
    deployment = os.environ.get("AZURE_OPENAI_EMBEDDING_DEPLOYMENT", "")
    vectors: list[list[float]] = []
    for i in range(0, len(texts), batch_size):
        resp = client.embeddings.create(model=deployment, input=texts[i:i + batch_size])
        vectors.extend(d.embedding for d in resp.data)
    arr = np.array(vectors, dtype=np.float32)
    arr /= np.linalg.norm(arr, axis=1, keepdims=True) + 1e-12  # normalise for cosine similarity
    return arr


def embedding_text(c: Chunk) -> str:
    # Including title and section helps search match questions phrased around topics.
    return f"{c.title} — {c.section}\n{c.text}"


def build_index(docs_dir: Path = DOCS_DIR, index_dir: Path = INDEX_DIR) -> int:
    chunks = load_chunks(docs_dir)
    if not chunks:
        raise RuntimeError(f"No documents found in {docs_dir}")
    vectors = embed([embedding_text(c) for c in chunks])
    index_dir.mkdir(parents=True, exist_ok=True)
    np.save(index_dir / "vectors.npy", vectors)
    (index_dir / "chunks.json").write_text(
        json.dumps([asdict(c) for c in chunks], indent=1), encoding="utf-8"
    )
    return len(chunks)


class Index:
    def __init__(self, index_dir: Path = INDEX_DIR):
        if not (index_dir / "vectors.npy").exists():
            raise RuntimeError("No index found. Run:  python ingest.py")
        self.vectors = np.load(index_dir / "vectors.npy")
        self.chunks = [Chunk(**c) for c in json.loads((index_dir / "chunks.json").read_text(encoding="utf-8"))]

    def search(self, query: str, k: int = TOP_K) -> list[tuple[Chunk, float]]:
        q = embed([query])[0]
        scores = self.vectors @ q
        order = np.argsort(-scores)[:k]
        return [(self.chunks[i], float(scores[i])) for i in order if scores[i] >= MIN_SCORE]


# --------------------------------------------------------------------------- #
# Answering
# --------------------------------------------------------------------------- #
def format_context(hits: list[tuple[Chunk, float]]) -> str:
    return "\n\n".join(
        f"[{n}] {c.title} — {c.section} (file: {c.source})\n{c.text}"
        for n, (c, _) in enumerate(hits, start=1)
    )


def answer(question: str, index: Index, history: list[dict] | None = None) -> tuple[str, list[tuple[Chunk, float]]]:
    """Return (answer_text, hits). `history` is prior chat turns [{role, content}, ...]."""
    hits = index.search(question)
    if not hits:
        return "I couldn't find that in the policy documents.", []

    messages = [{"role": "system", "content": SYSTEM_PROMPT}]
    for turn in (history or [])[-6:]:  # keep a little conversational memory
        messages.append({"role": turn["role"], "content": turn["content"]})
    messages.append({
        "role": "user",
        "content": f"Policy excerpts:\n\n{format_context(hits)}\n\nQuestion: {question}",
    })

    resp = get_client().chat.completions.create(
        model=os.environ.get("AZURE_OPENAI_CHAT_DEPLOYMENT", ""),
        messages=messages,
    )
    return resp.choices[0].message.content or "", hits
